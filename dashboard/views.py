import io
import json
import jdatetime
from datetime import datetime, timedelta
from decimal import Decimal

from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User as AuthUser
from django.db.models import (
    Q, Sum, Count, F, Avg, Min, Max,
    Case, When, Value, DecimalField,
)
from django.http import HttpResponse
from django.utils import timezone

from accounts.permissions import require_permission
from customers.models import Customer
from calls.models import Call
from products.models import Product
from sales.models import Invoice, InvoiceItem
from hr.models import Employee, Task, LeaveRequest


# فیلد Decimal استاندارد برای Case/When روی Invoice.total
DEC_FIELD = DecimalField(max_digits=15, decimal_places=0)
ZERO = Decimal('0')


# ═══════════════════════════════════════════════════════════
# توابع کمکی — سطح دسترسی و فیلتر کاربر
# ═══════════════════════════════════════════════════════════

def _is_manager(user):
    """آیا این کاربر می‌تواند همه‌ی داده‌ها را ببیند؟"""
    if not user or not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    return user.groups.filter(name='مدیر فروش').exists()


def _get_data_owner(user, filter_user_id=None):
    """
    تعیین «مالک داده» برای فیلتر:
    - مدیر بدون فیلتر → None (همه‌ی داده‌ها)
    - مدیر با فیلتر → کاربر انتخاب‌شده
    - غیرمدیر → خود کاربر
    """
    if _is_manager(user):
        if filter_user_id:
            try:
                return AuthUser.objects.get(pk=int(filter_user_id))
            except (AuthUser.DoesNotExist, ValueError, TypeError):
                return None
        return None  # همه
    return user


def _filtered_customers(owner):
    """مشتریان قابل مشاهده"""
    qs = Customer.objects.all()
    if owner is not None:
        qs = qs.filter(Q(assigned_to=owner) | Q(created_by=owner))
    return qs


def _filtered_calls(owner):
    """تماس‌های قابل مشاهده"""
    qs = Call.objects.all()
    if owner is not None:
        qs = qs.filter(Q(called_by=owner) | Q(created_by=owner))
    return qs


def _filtered_invoices(owner, statuses=None):
    """فاکتورهای قابل مشاهده"""
    if statuses is None:
        statuses = ['issued', 'paid']
    qs = Invoice.objects.filter(status__in=statuses)
    if owner is not None:
        qs = qs.filter(created_by=owner)
    return qs


# ═══════════════════════════════════════════════════════════
# توابع کمکی تاریخ
# ═══════════════════════════════════════════════════════════

def _to_jalali_month(g_date):
    try:
        j = jdatetime.date.fromgregorian(date=g_date)
        return f"{j.year}/{j.month:02d}"
    except Exception:
        return str(g_date)


def _get_month_range(year, month):
    try:
        start_j = jdatetime.date(year, month, 1)
        start_g = start_j.togregorian()
        if month == 12:
            end_j = jdatetime.date(year + 1, 1, 1)
        else:
            end_j = jdatetime.date(year, month + 1, 1)
        end_g = end_j.togregorian() - timedelta(days=1)
        return start_g, end_g
    except Exception:
        return None, None


def _parse_jalali_date(value):
    if not value:
        return None
    try:
        persian = '۰۱۲۳۴۵۶۷۸۹'
        arabic = '٠١٢٣٤٥٦٧٨٩'
        for i, (p, a) in enumerate(zip(persian, arabic)):
            value = value.replace(p, str(i)).replace(a, str(i))
        value = value.replace('-', '/').replace('.', '/')
        if ' ' in value:
            value = value.split(' ', 1)[0]
        y, m, d = map(int, value.split('/'))
        return jdatetime.date(y, m, d).togregorian()
    except (ValueError, AttributeError):
        return None


def _get_employee_choices():
    """کاربران فعال برای dropdown فیلتر"""
    return AuthUser.objects.filter(
        is_active=True,
        is_staff=True,
    ).order_by('first_name', 'last_name', 'username')


# ═══════════════════════════════════════════════════════════
# جمع‌آوری داده‌های داشبورد اصلی
# ═══════════════════════════════════════════════════════════

def _get_dashboard_data(user, from_date=None, to_date=None, filter_user_id=None):
    """جمع‌آوری داده‌های داشبورد با فیلتر تاریخ، نقش و کاربر"""
    today = timezone.now().date()
    now_j = jdatetime.date.fromgregorian(date=today)
    is_mgr = _is_manager(user)

    # ═══ تعیین مالک داده ═══
    owner = _get_data_owner(user, filter_user_id)
    filter_user_selected = None
    if is_mgr and owner is not None:
        filter_user_selected = owner

    # ═══ QuerySetها ═══
    invoices_qs = _filtered_invoices(owner)
    calls_qs = _filtered_calls(owner)
    customers_qs = _filtered_customers(owner)

    # ═══ فیلتر تاریخ ═══
    if from_date:
        invoices_qs = invoices_qs.filter(issue_date__gte=from_date)
        calls_qs = calls_qs.filter(call_date__date__gte=from_date)
        customers_qs = customers_qs.filter(created_at__date__gte=from_date)

    if to_date:
        invoices_qs = invoices_qs.filter(issue_date__lte=to_date)
        calls_qs = calls_qs.filter(call_date__date__lte=to_date)
        customers_qs = customers_qs.filter(created_at__date__lte=to_date)

    # ═══ KPI ها ═══
    total_customers = customers_qs.count()
    active_customers = customers_qs.filter(status='active').count()
    lead_customers = customers_qs.filter(status='lead').count()

    total_calls = calls_qs.count()

    # ماه جاری شمسی
    this_month_start, this_month_end = _get_month_range(now_j.year, now_j.month)
    calls_this_month = _filtered_calls(owner).filter(
        call_date__date__gte=this_month_start,
        call_date__date__lte=this_month_end,
    ).count()

    total_invoices = invoices_qs.count()
    total_sales = invoices_qs.aggregate(s=Sum('total'))['s'] or ZERO

    # ═══ بهینه‌سازی: فروش این ماه / ماه قبل / سال قبل (۱ کوئری) ═══
    month_start, month_end = _get_month_range(now_j.year, now_j.month)
    prev_year = now_j.year if now_j.month > 1 else now_j.year - 1
    prev_month = now_j.month - 1 if now_j.month > 1 else 12
    prev_start, prev_end = _get_month_range(prev_year, prev_month)
    last_year_start, last_year_end = _get_month_range(now_j.year - 1, now_j.month)

    comp = _filtered_invoices(owner).aggregate(
        this_m=Sum(Case(
            When(issue_date__gte=month_start, issue_date__lte=month_end, then=F('total')),
            default=Value(ZERO), output_field=DEC_FIELD,
        )),
        last_m=Sum(Case(
            When(issue_date__gte=prev_start, issue_date__lte=prev_end, then=F('total')),
            default=Value(ZERO), output_field=DEC_FIELD,
        )),
        last_y=Sum(Case(
            When(issue_date__gte=last_year_start, issue_date__lte=last_year_end, then=F('total')),
            default=Value(ZERO), output_field=DEC_FIELD,
        )),
    )

    if from_date or to_date:
        sales_this_month = total_sales
    else:
        sales_this_month = comp['this_m'] or ZERO

    sales_last_month = comp['last_m'] or ZERO
    sales_last_year = comp['last_y'] or ZERO

    # تغییرات
    if sales_last_month and sales_last_month > 0:
        sales_change = round(
            (float(sales_this_month) - float(sales_last_month))
            / float(sales_last_month) * 100, 1
        )
    else:
        sales_change = 100 if sales_this_month > 0 else 0

    if sales_last_year and sales_last_year > 0:
        yearly_change = round(
            (float(sales_this_month) - float(sales_last_year))
            / float(sales_last_year) * 100, 1
        )
    else:
        yearly_change = 100 if sales_this_month > 0 else 0

    # ═══ بهینه‌سازی: نمودار فروش ۱۲ ماه (۱ کوئری) ═══
    month_ranges = []
    for i in range(11, -1, -1):
        m = now_j.month - i
        y = now_j.year
        while m <= 0:
            m += 12
            y -= 1

        s_start, s_end = _get_month_range(y, m)
        if s_start is None:
            continue
        if from_date and s_end < from_date:
            continue
        if to_date and s_start > to_date:
            continue
        month_ranges.append((y, m, s_start, s_end))

    if month_ranges:
        when_kwargs = {}
        for idx, (_, _, s_start, s_end) in enumerate(month_ranges):
            when_kwargs[f'm{idx}'] = Sum(Case(
                When(issue_date__gte=s_start, issue_date__lte=s_end, then=F('total')),
                default=Value(ZERO), output_field=DEC_FIELD,
            ))
        result = _filtered_invoices(owner).aggregate(**when_kwargs)
        sales_data = [float(result[f'm{idx}'] or ZERO) for idx in range(len(month_ranges))]
        sales_labels = [f"{y}/{m:02d}" for (y, m, _, _) in month_ranges]
    else:
        sales_data = [0]
        sales_labels = ['بدون داده']

    # ═══ توزیع منابع جذب مشتری ═══
    sources_qs = customers_qs.values('source').annotate(
        count=Count('id')
    ).order_by('-count')

    source_labels_map = dict(Customer.SOURCE_CHOICES)
    source_data = []
    source_labels = []
    for item in sources_qs:
        source_labels.append(source_labels_map.get(item['source'], 'سایر'))
        source_data.append(item['count'])

    if not source_data:
        source_data = [0]
        source_labels = ['بدون داده']

    # ═══ توزیع وضعیت مشتریان ═══
    status_qs = customers_qs.values('status').annotate(
        count=Count('id')
    ).order_by('-count')

    status_labels_map = dict(Customer.STATUS_CHOICES)
    status_data = []
    status_labels = []
    for item in status_qs:
        status_labels.append(status_labels_map.get(item['status'], 'سایر'))
        status_data.append(item['count'])

    if not status_data:
        status_data = [0]
        status_labels = ['بدون داده']

    # ═══ توزیع وضعیت تماس‌ها ═══
    call_status_qs = calls_qs.values('status').annotate(
        count=Count('id')
    ).order_by('-count')

    call_status_map = dict(Call.CALL_STATUS_CHOICES)
    call_status_data = []
    call_status_labels = []
    for item in call_status_qs:
        call_status_labels.append(call_status_map.get(item['status'], 'سایر'))
        call_status_data.append(item['count'])

    if not call_status_data:
        call_status_data = [0]
        call_status_labels = ['بدون داده']

    # ═══ پرفروش‌ترین کالاها ═══
    top_products = InvoiceItem.objects.filter(
        invoice__in=invoices_qs
    ).values(name=F('product__name')).annotate(
        total_qty=Sum('quantity'),
        total_revenue=Sum('total_price')
    ).order_by('-total_qty')[:5]

    top_products_list = [
        {
            'name': p['name'],
            'qty': p['total_qty'],
            'revenue': float(p['total_revenue'] or 0),
        }
        for p in top_products
    ]

    # ═══ عملکرد کارمندان (فقط مدیرها، بدون فیلتر) ═══
    if is_mgr and filter_user_selected is None:
        employee_perf = invoices_qs.filter(
            created_by__isnull=False
        ).values(
            username=F('created_by__username'),
            first_name=F('created_by__first_name'),
            last_name=F('created_by__last_name'),
        ).annotate(
            total_sales=Sum('total'),
            invoice_count=Count('id')
        ).order_by('-total_sales')

        employee_perf_list = []
        for e in employee_perf:
            name = f"{e['first_name'] or ''} {e['last_name'] or ''}".strip()
            if not name:
                name = e['username']
            employee_perf_list.append({
                'name': name,
                'username': e['username'],
                'total_sales': float(e['total_sales'] or 0),
                'invoice_count': e['invoice_count'],
            })
    else:
        employee_perf_list = []

    emp_chart_names = [e['name'] for e in employee_perf_list[:10]]
    emp_chart_values = [e['total_sales'] for e in employee_perf_list[:10]]

    # ═══ قیف فروش ═══
    leads_qs = customers_qs.filter(status='lead')
    leads_count = leads_qs.count()

    leads_with_calls = _filtered_calls(owner).filter(
        customer__in=leads_qs
    ).values('customer').distinct().count()

    inv_funnel = _filtered_invoices(owner).filter(
        customer__in=leads_qs
    ).aggregate(
        with_invoice=Count('customer', distinct=True),
        paid=Count('customer', distinct=True, filter=Q(status='paid')),
    )
    leads_with_invoice = inv_funnel['with_invoice'] or 0
    paid_leads = inv_funnel['paid'] or 0

    funnel_data = {
        'labels': ['لید جدید', 'تماس گرفته', 'فاکتور صادر', 'پرداخت شده'],
        'values': [leads_count, leads_with_calls, leads_with_invoice, paid_leads],
    }

    # ═══ مقایسه فروش ═══
    if from_date or to_date:
        if from_date and to_date:
            range_days = (to_date - from_date).days + 1
            prev_range_end = from_date - timedelta(days=1)
            prev_range_start = prev_range_end - timedelta(days=range_days - 1)

            prev_sales = _filtered_invoices(owner).filter(
                issue_date__gte=prev_range_start,
                issue_date__lte=prev_range_end,
            ).aggregate(s=Sum('total'))['s'] or 0

            compare_data = {
                'labels': ['بازه قبل', 'بازه انتخابی'],
                'values': [float(prev_sales), float(total_sales)],
            }
        else:
            compare_data = {
                'labels': ['بازه انتخابی'],
                'values': [float(total_sales)],
            }
    else:
        compare_data = {
            'labels': ['ماه قبل', 'این ماه', 'سال قبل (همین ماه)'],
            'values': [
                float(sales_last_month),
                float(sales_this_month),
                float(sales_last_year),
            ],
        }

    # ═══ وظایف و مرخصی ═══
    if owner is not None:
        open_tasks = Task.objects.filter(
            status__in=['pending', 'in_progress'],
            assigned_to__user=owner
        ).count()
        overdue_tasks = Task.objects.filter(
            due_date__lt=today,
            assigned_to__user=owner
        ).exclude(status__in=['done', 'cancelled']).count()
        pending_leaves = LeaveRequest.objects.filter(
            status='pending',
            employee__user=owner
        ).count()
    else:
        open_tasks = Task.objects.filter(status__in=['pending', 'in_progress']).count()
        overdue_tasks = Task.objects.filter(
            due_date__lt=today
        ).exclude(status__in=['done', 'cancelled']).count()
        pending_leaves = LeaveRequest.objects.filter(status='pending').count()

    # ═══ لیست کاربران برای dropdown فیلتر (فقط مدیرها) ═══
    filter_users_list = []
    if is_mgr:
        filter_users_list = _get_employee_choices()

    # ═══ آماده‌سازی JSON ═══
    def _json(obj):
        return json.dumps(obj, ensure_ascii=False)

    return {
        'today': now_j.strftime('%Y/%m/%d'),
        'is_manager_view': is_mgr,
        'filter_user_selected': filter_user_selected,
        'filter_users_list': filter_users_list,

        'kpis': {
            'total_customers': total_customers,
            'active_customers': active_customers,
            'lead_customers': lead_customers,
            'total_calls': total_calls,
            'calls_this_month': calls_this_month,
            'total_invoices': total_invoices,
            'total_sales': total_sales,
            'sales_this_month': sales_this_month,
            'sales_last_month': sales_last_month,
            'sales_last_year': sales_last_year,
            'sales_change': sales_change,
            'yearly_change': yearly_change,
            'open_tasks': open_tasks,
            'overdue_tasks': overdue_tasks,
            'pending_leaves': pending_leaves,
        },

        'sales_data': sales_data,
        'sales_labels': sales_labels,
        'source_data': source_data,
        'source_labels': source_labels,

        'sales_data_json': _json(sales_data),
        'sales_labels_json': _json(sales_labels),
        'source_data_json': _json(source_data),
        'source_labels_json': _json(source_labels),
        'status_data_json': _json(status_data),
        'status_labels_json': _json(status_labels),
        'call_status_data_json': _json(call_status_data),
        'call_status_labels_json': _json(call_status_labels),
        'emp_chart_names_json': _json(emp_chart_names),
        'emp_chart_values_json': _json(emp_chart_values),

        'compare_data': compare_data,
        'funnel': funnel_data,
        'top_products': top_products_list,
        'employee_perf': employee_perf_list,
    }


# ═══════════════════════════════════════════════════════════
# Views
# ═══════════════════════════════════════════════════════════

@login_required
@require_permission('dashboard.view')
def home(request):
    """داشبورد اصلی با فیلتر تاریخ و کاربر"""
    from_date_str = request.GET.get('from_date', '').strip()
    to_date_str = request.GET.get('to_date', '').strip()
    filter_user_id = request.GET.get('employee', '').strip()

    from_date = _parse_jalali_date(from_date_str) if from_date_str else None
    to_date = _parse_jalali_date(to_date_str) if to_date_str else None

    context = _get_dashboard_data(
        request.user, from_date, to_date,
        filter_user_id=filter_user_id or None,
    )
    context['from_date_jalali'] = from_date_str
    context['to_date_jalali'] = to_date_str
    context['filter_user_id'] = filter_user_id

    return render(request, 'dashboard/home.html', context)


@login_required
def global_search(request):
    """جستجوی سراسری — فیلترشده بر اساس نقش"""
    query = request.GET.get('q', '').strip()
    results = {'customers': [], 'calls': [], 'products': [], 'total': 0}

    owner = _get_data_owner(request.user)

    if query and len(query) >= 2:
        results['customers'] = list(_filtered_customers(owner).filter(
            Q(name__icontains=query) | Q(mobile__icontains=query) |
            Q(email__icontains=query) | Q(national_id__icontains=query)
        )[:10])

        results['calls'] = list(_filtered_calls(owner).filter(
            Q(caller_name__icontains=query) | Q(call_number__icontains=query) |
            Q(notes__icontains=query)
        )[:10])

        results['products'] = list(Product.objects.filter(
            Q(name__icontains=query) | Q(code__icontains=query) |
            Q(description__icontains=query)
        )[:10])

        results['total'] = (len(results['customers']) + len(results['calls'])
                            + len(results['products']))

    return render(request, 'dashboard/search.html', {
        'query': query, 'results': results,
    })


@login_required
@require_permission('dashboard.export')
def export_excel(request):
    """خروجی Excel — فیلترشده بر اساس نقش و کاربر"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        return HttpResponse(
            'openpyxl نصب نیست. لطفاً pip install openpyxl را اجرا کنید.',
            status=500
        )

    from_date_str = request.GET.get('from_date', '').strip()
    to_date_str = request.GET.get('to_date', '').strip()
    filter_user_id = request.GET.get('employee', '').strip()

    from_date = _parse_jalali_date(from_date_str) if from_date_str else None
    to_date = _parse_jalali_date(to_date_str) if to_date_str else None

    data = _get_dashboard_data(
        request.user, from_date, to_date,
        filter_user_id=filter_user_id or None,
    )

    wb = Workbook()

    header_font = Font(bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill('solid', fgColor='2E86AB')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    # ═══ Sheet 1: خلاصه ═══
    ws = wb.active
    ws.title = 'خلاصه'
    ws.sheet_view.rightToLeft = True

    ws.append(['گزارش CRM دیجیران'])
    ws.merge_cells('A1:C1')
    ws['A1'].font = Font(bold=True, size=16, color='2E86AB')
    ws['A1'].alignment = center
    ws.row_dimensions[1].height = 30

    ws.append([])
    ws.append([f"تاریخ گزارش: {data['today']}"])
    if data.get('filter_user_selected'):
        u = data['filter_user_selected']
        ws.append([f"فیلتر کاربر: {u.get_full_name() or u.username}"])
    elif not data.get('is_manager_view'):
        ws.append(['⚠️ این گزارش فقط شامل داده‌های شماست'])
    ws.append([])

    ws.append(['شاخص', 'مقدار', 'واحد'])
    for cell in ws[ws.max_row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    kpis = data['kpis']
    kpi_rows = [
        ('کل مشتریان', kpis['total_customers'], 'نفر'),
        ('مشتریان فعال', kpis['active_customers'], 'نفر'),
        ('مشتریان بالقوه (لید)', kpis['lead_customers'], 'نفر'),
        ('کل تماس‌ها', kpis['total_calls'], 'تماس'),
        ('تماس‌های این ماه', kpis['calls_this_month'], 'تماس'),
        ('کل فاکتورها', kpis['total_invoices'], 'فاکتور'),
        ('مجموع فروش', int(kpis['total_sales']), 'تومان'),
        ('فروش این ماه', int(kpis['sales_this_month']), 'تومان'),
        ('فروش ماه قبل', int(kpis['sales_last_month']), 'تومان'),
        ('تغییر ماهانه', f"{kpis['sales_change']}٪", '—'),
        ('فروش سال قبل (همین ماه)', int(kpis['sales_last_year']), 'تومان'),
        ('تغییر سالانه', f"{kpis['yearly_change']}٪", '—'),
        ('وظایف باز', kpis['open_tasks'], 'وظیفه'),
        ('وظایف عقب‌افتاده', kpis['overdue_tasks'], 'وظیفه'),
        ('مرخصی‌های در انتظار', kpis['pending_leaves'], 'درخواست'),
    ]

    for row in kpi_rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.alignment = center
            cell.border = border

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 15

    # ═══ Sheet 2: فروش ماهانه ═══
    ws2 = wb.create_sheet('فروش ماهانه')
    ws2.sheet_view.rightToLeft = True

    ws2.append(['ماه', 'فروش (تومان)'])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for label, value in zip(data['sales_labels'], data['sales_data']):
        ws2.append([label, int(value)])
        for cell in ws2[ws2.max_row]:
            cell.alignment = center
            cell.border = border

    ws2.column_dimensions['A'].width = 20
    ws2.column_dimensions['B'].width = 25

    # ═══ Sheet 3: عملکرد کارمندان ═══
    if data.get('employee_perf'):
        ws3 = wb.create_sheet('عملکرد کارمندان')
        ws3.sheet_view.rightToLeft = True

        ws3.append(['کارمند', 'تعداد فاکتور', 'مجموع فروش (تومان)'])
        for cell in ws3[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        for emp in data['employee_perf']:
            ws3.append([emp['name'], emp['invoice_count'], int(emp['total_sales'])])
            for cell in ws3[ws3.max_row]:
                cell.alignment = center
                cell.border = border

        ws3.column_dimensions['A'].width = 25
        ws3.column_dimensions['B'].width = 15
        ws3.column_dimensions['C'].width = 25

    # ═══ Sheet 4: پرفروش‌ترین کالاها ═══
    ws4 = wb.create_sheet('پرفروش‌ترین کالاها')
    ws4.sheet_view.rightToLeft = True

    ws4.append(['کالا', 'تعداد فروش', 'درآمد (تومان)'])
    for cell in ws4[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for p in data['top_products']:
        ws4.append([p['name'], p['qty'], int(p['revenue'])])
        for cell in ws4[ws4.max_row]:
            cell.alignment = center
            cell.border = border

    ws4.column_dimensions['A'].width = 30
    ws4.column_dimensions['B'].width = 15
    ws4.column_dimensions['C'].width = 25

    # ═══ Sheet 5: منابع جذب ═══
    ws5 = wb.create_sheet('منابع جذب')
    ws5.sheet_view.rightToLeft = True

    ws5.append(['منبع', 'تعداد مشتری'])
    for cell in ws5[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for label, value in zip(data['source_labels'], data['source_data']):
        ws5.append([label, value])
        for cell in ws5[ws5.max_row]:
            cell.alignment = center
            cell.border = border

    ws5.column_dimensions['A'].width = 25
    ws5.column_dimensions['B'].width = 15

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'crm-report-{datetime.now().strftime("%Y%m%d-%H%M")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ═══════════════════════════════════════════════════════════
# BI Dashboard — فاز ۷.۷
# ═══════════════════════════════════════════════════════════

def _bi_heatmap_data(calls_qs):
    """
    Heatmap ساعت‌های طلایی تماس‌ها.
    Returns: (matrix[7][24], day_labels, hour_labels)
    """
    # ترتیب روزهای هفته در ایران: شنبه ... جمعه
    # weekday() پایتون: 0=Monday, 5=Saturday, 6=Sunday
    day_order = [5, 6, 0, 1, 2, 3, 4]
    day_labels = ['شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه']

    matrix = [[0] * 24 for _ in range(7)]
    for call_date, _status in calls_qs.values_list('call_date', 'status'):
        if call_date is None:
            continue
        py_dow = call_date.weekday()
        try:
            row = day_order.index(py_dow)
        except ValueError:
            continue
        matrix[row][call_date.hour] += 1

    hour_labels = [f'{h:02d}' for h in range(24)]
    return matrix, day_labels, hour_labels


@login_required
@require_permission('dashboard.view')
def bi_dashboard(request):
    """صفحه‌ی BI — گزارش‌های عمیق با نمودارهای تعاملی"""
    today = timezone.now().date()
    now_j = jdatetime.date.fromgregorian(date=today)
    is_mgr = _is_manager(request.user)

    # ─── فیلتر کاربر ───
    filter_user_id = request.GET.get('employee', '').strip() or None
    owner = _get_data_owner(request.user, filter_user_id)
    filter_user_selected = None
    if is_mgr and owner is not None:
        filter_user_selected = owner

    invoices_qs = _filtered_invoices(owner)
    calls_qs = _filtered_calls(owner)
    customers_qs = _filtered_customers(owner)

    # ═══════════════════════════════════════════
    # KPI ها
    # ═══════════════════════════════════════════
    total_customers = customers_qs.count()
    active_customers = customers_qs.filter(status='active').count()
    lead_customers = customers_qs.filter(status='lead').count()

    conversion_rate = round(
        (active_customers / total_customers * 100) if total_customers else 0, 1
    )

    invoice_stats = invoices_qs.aggregate(
        total_amount=Sum('total'),
        count=Count('id'),
    )
    avg_invoice = (
        (invoice_stats['total_amount'] or 0) / invoice_stats['count']
        if invoice_stats['count'] else 0
    )

    customers_with_first_invoice = (
        customers_qs
        .filter(invoices__status__in=['issued', 'paid'])
        .annotate(first_invoice_date=Min('invoices__issue_date'))
        .filter(first_invoice_date__isnull=False)
    )
    total_days = 0
    count_days = 0
    for c in customers_with_first_invoice.only('first_contact_date'):
        first_inv = getattr(c, 'first_invoice_date', None)
        if first_inv and c.first_contact_date:
            delta = (first_inv - c.first_contact_date.date()).days
            if delta >= 0:
                total_days += delta
                count_days += 1
    avg_conversion_days = round(total_days / count_days, 1) if count_days else 0

    # ═══════════════════════════════════════════
    # Trend فروش ۱۲ ماه
    # ═══════════════════════════════════════════
    trend_labels = []
    trend_current = []
    trend_previous = []
    for i in range(11, -1, -1):
        m = now_j.month - i
        y = now_j.year
        while m <= 0:
            m += 12
            y -= 1
        s_start, s_end = _get_month_range(y, m)
        if s_start is None:
            continue

        cur = invoices_qs.filter(
            issue_date__gte=s_start, issue_date__lte=s_end
        ).aggregate(s=Sum('total'))['s'] or 0
        trend_current.append(float(cur))
        trend_labels.append(f'{y}/{m:02d}')

        prev_start, prev_end = _get_month_range(y - 1, m)
        if prev_start:
            prev = invoices_qs.filter(
                issue_date__gte=prev_start, issue_date__lte=prev_end
            ).aggregate(s=Sum('total'))['s'] or 0
            trend_previous.append(float(prev))
        else:
            trend_previous.append(0)

    # ═══════════════════════════════════════════
    # Heatmap ساعت‌های طلایی
    # ═══════════════════════════════════════════
    heatmap_matrix, heatmap_days, heatmap_hours = _bi_heatmap_data(calls_qs)
    heatmap_series = []
    for i, day in enumerate(heatmap_days):
        heatmap_series.append({
            'name': day,
            'data': [
                {'x': heatmap_hours[h], 'y': heatmap_matrix[i][h]}
                for h in range(24)
            ],
        })

    # ═══════════════════════════════════════════
    # نرخ تبدیل هر منبع جذب
    # ═══════════════════════════════════════════
    source_map = dict(Customer.SOURCE_CHOICES)
    sources_qs = customers_qs.values('source').annotate(
        total=Count('id'),
        active=Count('id', filter=Q(status='active')),
    ).order_by('-total')

    source_data = []
    for s in sources_qs:
        src_key = s['source']
        total = s['total']
        active = s['active']
        rate = round((active / total * 100) if total else 0, 1)
        source_data.append({
            'label': source_map.get(src_key, 'سایر'),
            'total': total,
            'active': active,
            'rate': rate,
        })

    # ═══════════════════════════════════════════
    # Top 10 مشتریان
    # ═══════════════════════════════════════════
    top_customers = (
        customers_qs
        .annotate(
            total_purchase=Sum(
                'invoices__total',
                filter=Q(invoices__status__in=['issued', 'paid']),
            ),
            invoice_count=Count(
                'invoices',
                filter=Q(invoices__status__in=['issued', 'paid']),
            ),
        )
        .filter(total_purchase__gt=0)
        .order_by('-total_purchase')[:10]
    )

    # ═══════════════════════════════════════════
    # Top 10 کالاها
    # ═══════════════════════════════════════════
    top_products = (
        InvoiceItem.objects
        .filter(invoice__in=invoices_qs)
        .values('product__name', 'product__code')
        .annotate(
            total_qty=Sum('quantity'),
            total_revenue=Sum('total_price'),
        )
        .order_by('-total_revenue')[:10]
    )

    # ═══════════════════════════════════════════
    # Top 10 کارمندان
    # ═══════════════════════════════════════════
    top_employees = []
    if owner is None:
        emp_qs = (
            invoices_qs
            .filter(created_by__isnull=False)
            .values(
                'created_by__username',
                'created_by__first_name',
                'created_by__last_name',
            )
            .annotate(
                total_sales=Sum('total'),
                invoice_count=Count('id'),
            )
            .order_by('-total_sales')[:10]
        )
        for e in emp_qs:
            name = f"{e['created_by__first_name'] or ''} {e['created_by__last_name'] or ''}".strip()
            if not name:
                name = e['created_by__username']
            top_employees.append({
                'name': name,
                'username': e['created_by__username'],
                'total_sales': float(e['total_sales'] or 0),
                'invoice_count': e['invoice_count'],
            })

    # ═══════════════════════════════════════════
    # مشتریان غیرفعال (بدون خرید ۹۰+ روز)
    # ═══════════════════════════════════════════
    cutoff = today - timedelta(days=90)

    inactive_customers = (
        customers_qs
        .annotate(
            last_invoice=Max(
                'invoices__issue_date',
                filter=Q(invoices__status__in=['issued', 'paid']),
            ),
        )
        .filter(
            Q(last_invoice__lt=cutoff) | Q(last_invoice__isnull=True),
            status__in=['active', 'inactive'],
        )
        .order_by('last_invoice')[:10]
    )

    # ═══════════════════════════════════════════
    # لیست کاربران برای فیلتر
    # ═══════════════════════════════════════════
    filter_users_list = []
    if is_mgr:
        filter_users_list = _get_employee_choices()

    # ═══════════════════════════════════════════
    # آماده‌سازی JSON
    # ═══════════════════════════════════════════
    def _json(obj):
        return json.dumps(obj, ensure_ascii=False)

    context = {
        'today': now_j.strftime('%Y/%m/%d'),
        'is_manager_view': is_mgr,
        'filter_user_selected': filter_user_selected,
        'filter_users_list': filter_users_list,
        'filter_user_id': filter_user_id or '',

        'kpis': {
            'total_customers': total_customers,
            'active_customers': active_customers,
            'lead_customers': lead_customers,
            'conversion_rate': conversion_rate,
            'avg_invoice': float(avg_invoice),
            'avg_conversion_days': avg_conversion_days,
        },

        'trend_labels': trend_labels,
        'trend_current': trend_current,
        'trend_previous': trend_previous,

        'heatmap_series': heatmap_series,
        'heatmap_hours': heatmap_hours,

        'source_data': source_data,

        'top_customers': top_customers,
        'top_products': top_products,
        'top_employees': top_employees,
        'inactive_customers': inactive_customers,

        'trend_labels_json': _json(trend_labels),
        'trend_current_json': _json(trend_current),
        'trend_previous_json': _json(trend_previous),
        'heatmap_series_json': _json(heatmap_series),
        'heatmap_hours_json': _json(heatmap_hours),
        'source_data_json': _json(source_data),
    }
    return render(request, 'dashboard/bi.html', context)