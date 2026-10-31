from crm.numerals import digit_contains
import json
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods

import jdatetime

from .forms import AdForm, AdPlatformForm, AdProcessForm, AdSourceForm
from .models import Ad, AdLog, AdPlatform, AdProcess, AdSource
from accounts.permissions import has_permission as _has_permission

# ─────────────────────────────────────────────
# سیستم بررسی دسترسی یکپارچه
# ─────────────────────────────────────────────
def check_permission(user, perm_key):
    """بررسی دسترسی با استفاده از سیستم مرکزی RBAC"""
    if user.is_superuser:
        return True
    return _has_permission(user, perm_key)

def perm_required(perm_key):
    """دکوراتور برای بررسی دسترسی در ویوها"""
    def decorator(view_func):
        @login_required
        def _wrapped_view(request, *args, **kwargs):
            if not check_permission(request.user, perm_key):
                return HttpResponseForbidden('شما دسترسی لازم برای این عملیات را ندارید.')
            return view_func(request, *args, **kwargs)
        return _wrapped_view
    return decorator


# ─────────────────────────────────────────────
# Helper: تبدیل تاریخ شمسی به میلادی
# ─────────────────────────────────────────────
def _parse_jalali(s):
    if not s:
        return None
    try:
        parts = s.strip().replace('-', '/').split('/')
        if len(parts) != 3:
            return None
        y, m, d = map(int, parts)
        return jdatetime.date(y, m, d).togregorian()
    except (ValueError, TypeError):
        return None


def _apply_report_filters(request, qs):
    report_user_id = request.GET.get('report_user', '')
    report_from = request.GET.get('report_from', '')
    report_to = request.GET.get('report_to', '')

    if request.user.is_superuser and report_user_id:
        qs = qs.filter(created_by_id=report_user_id)

    g_from = _parse_jalali(report_from)
    if g_from:
        qs = qs.filter(created_at__date__gte=g_from)

    g_to = _parse_jalali(report_to)
    if g_to:
        qs = qs.filter(created_at__date__lte=g_to)

    return qs, report_user_id, report_from, report_to


# ─────────────────────────────────────────────
# لیست تبلیغات + گزارش
# ─────────────────────────────────────────────
@login_required
@perm_required('ads.view')
def ad_list(request):
    qs = Ad.objects.select_related(
        'platform', 'process', 'source', 'product', 'created_by',
    )
    if not request.user.is_superuser:
        qs = qs.filter(created_by=request.user)

    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(
            digit_contains('ad_number', q)
            | digit_contains('platform__name', q)
            | digit_contains('process__name', q)
            | digit_contains('source_phone_snapshot', q)
            | digit_contains('source_label_snapshot', q)
            | digit_contains('product__name', q)
            | digit_contains('note', q)
        )

    platform_id = request.GET.get('platform')
    if platform_id:
        qs = qs.filter(platform_id=platform_id)

    if request.user.is_superuser:
        user_id = request.GET.get('user')
        if user_id:
            qs = qs.filter(created_by_id=user_id)

    now = timezone.now()
    stats = {
        'total': qs.count(),
        'total_cost': qs.aggregate(t=Sum('cost'))['t'] or 0,
        'this_month': qs.filter(
            created_at__year=now.year,
            created_at__month=now.month,
        ).count(),
        'platforms': AdPlatform.objects.filter(is_active=True).count(),
    }

    paginator = Paginator(qs, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    # گزارش
    report_qs = Ad.objects.all()
    if not request.user.is_superuser:
        report_qs = report_qs.filter(created_by=request.user)

    report_qs, r_user, r_from, r_to = _apply_report_filters(request, report_qs)

    report_summary = {
        'count': report_qs.count(),
        'total_cost': report_qs.aggregate(t=Sum('cost'))['t'] or 0,
    }

    by_platform = list(
        report_qs.values('platform__name')
        .annotate(count=Count('id'), total=Sum('cost'))
        .order_by('-total')
    )
    for item in by_platform:
        item['total'] = item['total'] or 0

    by_user = None
    if request.user.is_superuser:
        by_user = list(
            report_qs.values('created_by__username')
            .annotate(count=Count('id'), total=Sum('cost'))
            .order_by('-total')
        )
        for item in by_user:
            item['total'] = item['total'] or 0

    report = {
        'user_id': r_user,
        'from': r_from,
        'to': r_to,
        'summary': report_summary,
        'by_platform': by_platform,
        'by_user': by_user,
        'by_source': list(report_qs.values('source_phone_snapshot', 'source_label_snapshot').annotate(count=Count('id'), total=Sum('cost')).order_by('-total')),
        'chart_labels': [item['platform__name'] for item in by_platform],
        'chart_values': [float(item['total']) for item in by_platform],
    }

    # JSON ایمن برای بلوک <script> (جلوگیری از XSS نام پلتفرم/فرآیند)
    def _json(obj):
        s = json.dumps(obj, ensure_ascii=False)
        return (s.replace('<', '\\u003C')
                 .replace('>', '\\u003E')
                 .replace('&', '\\u0026'))

    report['chart_labels_json'] = _json(report.get('chart_labels') or [])
    report['chart_values_json'] = _json(report.get('chart_values') or [])

    context = {
        'page_obj': page_obj,
        'q': q,
        'platforms': AdPlatform.objects.filter(is_active=True),
        'current_platform': platform_id,
        'users': (
            Ad.objects.values_list('created_by__id', 'created_by__username')
            .distinct()
            if request.user.is_superuser else None
        ),
        'is_admin': request.user.is_superuser,
        'stats': stats,
        'report': report,
        'can_view': check_permission(request.user, 'ads.view'),
        'can_create': check_permission(request.user, 'ads.add'),
        'can_edit': check_permission(request.user, 'ads.edit'),
        'can_delete': check_permission(request.user, 'ads.delete'),
        'can_manage_definitions': check_permission(request.user, 'ads.manage_defs'),
    }
    return render(request, 'ads/ad_list.html', context)


# ─────────────────────────────────────────────
# خروجی Excel گزارش
# ─────────────────────────────────────────────
@login_required
@perm_required('ads.view')
def ad_export_report(request):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    qs = Ad.objects.select_related(
        'platform', 'process', 'source', 'product', 'created_by',
    ).order_by('-created_at')

    if not request.user.is_superuser:
        qs = qs.filter(created_by=request.user)

    qs, _, _, _ = _apply_report_filters(request, qs)

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = 'ریز تبلیغات'
    ws.sheet_view.rightToLeft = True

    headers = [
        'ردیف', 'عنوان', 'فرآیند', 'هزینه (تومان)', 'منبع',
        'محصول', 'شناسه آگهی', 'کارشناس', 'تاریخ ثبت', 'توضیحات',
    ]
    ws.append(headers)

    header_fill = PatternFill('solid', fgColor='2563EB')
    header_font = Font(color='FFFFFF', bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for i, ad in enumerate(qs, 1):
        ws.append([
            i,
            ad.platform.name,
            ad.process.name,
            float(ad.cost),
            ad.source_display,
            ad.product.name if ad.product else '—',
            ad.ad_number,
            ad.created_by.username,
            ad.created_at.strftime('%Y/%m/%d %H:%M'),
            ad.note or '',
        ])

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col_idx)].width = 18

    ws2 = wb.create_sheet('تفکیک پلتفرم')
    ws2.sheet_view.rightToLeft = True
    ws2.append(['عنوان', 'تعداد', 'جمع هزینه (تومان)'])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font

    by_platform = (
        qs.values('platform__name')
        .annotate(count=Count('id'), total=Sum('cost'))
        .order_by('-total')
    )
    for item in by_platform:
        ws2.append([
            item['platform__name'],
            item['count'],
            float(item['total'] or 0),
        ])

    if request.user.is_superuser:
        ws3 = wb.create_sheet('تفکیک کارشناس')
        ws3.sheet_view.rightToLeft = True
        ws3.append(['کارشناس', 'تعداد', 'جمع هزینه (تومان)'])
        for cell in ws3[1]:
            cell.fill = header_fill
            cell.font = header_font

        by_user = (
            qs.values('created_by__username')
            .annotate(count=Count('id'), total=Sum('cost'))
            .order_by('-total')
        )
        for item in by_user:
            ws3.append([
                item['created_by__username'],
                item['count'],
                float(item['total'] or 0),
            ])

    source_sheet = wb.create_sheet('تفکیک منبع هنگام ثبت')
    source_sheet.sheet_view.rightToLeft = True
    source_sheet.append(['شماره منبع هنگام ثبت', 'برچسب هنگام ثبت', 'تعداد', 'جمع هزینه (تومان)'])
    for cell in source_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for item in qs.values('source_phone_snapshot', 'source_label_snapshot').annotate(count=Count('id'), total=Sum('cost')).order_by('-total'):
        source_sheet.append([item['source_phone_snapshot'], item['source_label_snapshot'], item['count'], float(item['total'] or 0)])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'ads_report_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ─────────────────────────────────────────────
# ایجاد تبلیغ
# ─────────────────────────────────────────────
@login_required
@perm_required('ads.add')
@require_http_methods(['GET', 'POST'])
def ad_create(request):
    if request.method == 'POST':
        form = AdForm(request.POST)
        if form.is_valid():
            with transaction.atomic():
                ad = form.save(commit=False)
                ad.created_by = request.user
                ad.save()

                AdLog.objects.create(
                    ad=ad,
                    action='create',
                    changed_by=request.user,
                    changes={
                        'platform': [None, ad.platform.name],
                        'process': [None, ad.process.name],
                        'cost': [None, str(ad.cost)],
                        'source': [None, ad.source_display],
                        'product': [None, ad.product.name if ad.product else '—'],
                        'ad_number': [None, ad.ad_number],
                    },
                )
            messages.success(request, 'تبلیغ با موفقیت ثبت شد.')
            return redirect('ads:list')
    else:
        form = AdForm()

    return render(request, 'ads/ad_form.html', {
        'form': form,
        'title': 'ثبت تبلیغ جدید',
        'is_edit': False,
    })


# ─────────────────────────────────────────────
# ویرایش تبلیغ (و مشاهده)
# ─────────────────────────────────────────────
@login_required
@require_http_methods(['GET', 'POST'])
def ad_edit(request, pk):
    ad = get_object_or_404(Ad, pk=pk)

    # ⭐ Fix #12: ads.view برای GET و POST هر دو چک می‌شود (پیش‌نیاز منطقی)
    # قبلاً اگر کاربر override داشت (edit=True, view=False)، با POST مستقیم
    # (curl) می‌توانست داده را تغییر دهد بدون اینکه هرگز فرم را دیده باشد.
    if not check_permission(request.user, 'ads.view'):
        return HttpResponseForbidden('شما دسترسی مشاهده ندارید.')

    # POST علاوه بر view، نیاز به edit هم دارد
    if request.method == 'POST' and not check_permission(request.user, 'ads.edit'):
        return HttpResponseForbidden('شما دسترسی ویرایش ندارید.')

    # چک مالکیت رکورد (غیرسوپریوزر فقط رکوردهای خودش)
    if not request.user.is_superuser and ad.created_by_id != request.user.id:
        return HttpResponseForbidden('فقط رکوردهای خودتان')

    if request.method == 'POST':
        old = {
            'platform': ad.platform.name,
            'process': ad.process.name,
            'cost': str(ad.cost),
            'source': ad.source_display,
            'product': ad.product.name if ad.product else '—',
            'ad_number': ad.ad_number,
        }

        form = AdForm(request.POST, instance=ad)
        if form.is_valid():
            with transaction.atomic():
                updated = form.save()

                new_vals = {
                    'platform': updated.platform.name,
                    'process': updated.process.name,
                    'cost': str(updated.cost),
                    'source': updated.source_display,
                    'product': updated.product.name if updated.product else '—',
                    'ad_number': updated.ad_number,
                }

                changes = {
                    field: [old[field], new_vals[field]]
                    for field in old
                    if old[field] != new_vals[field]
                }

                if changes:
                    AdLog.objects.create(
                        ad=updated,
                        action='update',
                        changed_by=request.user,
                        changes=changes,
                    )

            messages.success(request, 'تبلیغ با موفقیت ویرایش شد.')
            return redirect('ads:list')
    else:
        form = AdForm(instance=ad)

    return render(request, 'ads/ad_form.html', {
        'form': form,
        'title': 'ویرایش تبلیغ',
        'is_edit': True,
        'ad': ad,
        'logs': ad.logs.select_related('changed_by').all()[:20],
    })


# ─────────────────────────────────────────────
# حذف تبلیغ
# ─────────────────────────────────────────────
@login_required
@perm_required('ads.delete')
@require_http_methods(['POST'])
def ad_delete(request, pk):
    ad = get_object_or_404(Ad, pk=pk)
    with transaction.atomic():
        # ⭐ Fix #2: AdLog.ad حالا SET_NULL است.
        # پس از ad.delete()، لاگ با ad_id=NULL باقی می‌ماند و
        # ad_number_snapshot هویت آگهی را حفظ می‌کند.
        AdLog.objects.create(
            ad=ad,
            action='delete',
            changed_by=request.user,
            note=f'حذف تبلیغ {ad.ad_number}',
        )
        ad.delete()
    messages.success(request, 'تبلیغ حذف شد. لاگ حذف در سیستم ثبت شد.')
    return redirect('ads:list')


# ─────────────────────────────────────────────
# مدیریت عنوان‌ها، فرآیندها و منابع
# ─────────────────────────────────────────────
@login_required
@perm_required('ads.manage_defs')
def platform_list(request):
    if request.method == 'POST':
        form = AdPlatformForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'عنوان اضافه شد.')
            return redirect('ads:platform_list')
    else:
        form = AdPlatformForm()
    return render(request, 'ads/platform_list.html', {
        'form': form, 'items': AdPlatform.objects.all(),
    })


@login_required


@perm_required('ads.manage_defs')
@require_http_methods(['POST'])
def platform_delete(request, pk):
    obj = get_object_or_404(AdPlatform, pk=pk)
    if obj.ads.exists():
        messages.error(request, 'این عنوان در تبلیغات استفاده شده و قابل حذف نیست.')
    else:
        obj.delete()
        messages.success(request, 'عنوان حذف شد.')
    return redirect('ads:platform_list')


@login_required


@perm_required('ads.manage_defs')
def process_list(request):
    if request.method == 'POST':
        form = AdProcessForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'فرآیند اضافه شد.')
            return redirect('ads:process_list')
    else:
        form = AdProcessForm()
    return render(request, 'ads/process_list.html', {
        'form': form,
        'items': AdProcess.objects.select_related('platform').all(),
    })


@login_required


@perm_required('ads.manage_defs')
@require_http_methods(['POST'])
def process_delete(request, pk):
    obj = get_object_or_404(AdProcess, pk=pk)
    if obj.ads.exists():
        messages.error(request, 'این فرآیند در تبلیغات استفاده شده و قابل حذف نیست.')
    else:
        obj.delete()
        messages.success(request, 'فرآیند حذف شد.')
    return redirect('ads:process_list')


@login_required


@perm_required('ads.manage_defs')
def source_list(request):
    if request.method == 'POST':
        form = AdSourceForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'منبع اضافه شد.')
            return redirect('ads:source_list')
    else:
        form = AdSourceForm()
    return render(request, 'ads/source_list.html', {
        'form': form, 'items': AdSource.objects.all(),
    })


@login_required
@perm_required('ads.manage_defs')
@require_http_methods(['GET', 'POST'])
def source_edit(request, pk):
    source = get_object_or_404(AdSource, pk=pk)
    form = AdSourceForm(request.POST or None, instance=source)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'منبع به‌روزرسانی شد؛ شماره و برچسب ثبت‌شده در سوابق قبلی حفظ شدند.')
        return redirect('ads:source_list')
    return render(request, 'ads/source_list.html', {
        'form': form, 'items': AdSource.objects.all(), 'editing_source': source,
    })


@login_required


@perm_required('ads.manage_defs')
@require_http_methods(['POST'])
def source_delete(request, pk):
    obj = get_object_or_404(AdSource, pk=pk)
    if obj.ads.exists() or obj.calls.exists() or obj.customer_set.exists():
        messages.error(request, 'این منبع در سوابق استفاده شده است؛ برای توقف استفاده، آن را غیرفعال کنید.')
    else:
        obj.delete()
        messages.success(request, 'منبع حذف شد.')
    return redirect('ads:source_list')


# ─────────────────────────────────────────────
# API ها
# ─────────────────────────────────────────────
@login_required
@perm_required('ads.view')
def platform_processes_api(request, pk):
    processes = AdProcess.objects.filter(
        platform_id=pk, is_active=True,
    ).values('id', 'name', 'default_cost')
    return JsonResponse({'processes': list(processes)})


@login_required


@perm_required('ads.view')
def process_cost_api(request, pk):
    process = get_object_or_404(AdProcess, pk=pk, is_active=True)
    return JsonResponse({
        'cost': str(process.default_cost),
        'name': process.name,
        'platform_id': process.platform_id,
    })


@login_required


@perm_required('ads.view')
def platform_next_ad_number_api(request, pk):
    platform = get_object_or_404(AdPlatform, pk=pk, is_active=True)
    return JsonResponse({
        'ad_number': platform.get_next_ad_number(),
        'platform_name': platform.name,
    })
