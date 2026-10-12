import json
import jdatetime
from datetime import timedelta

from django.contrib.auth.models import User
from django.core.paginator import Paginator
from django.db.models import Q, Count
from django.http import HttpResponse
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.utils import timezone

from accounts.permissions import require_permission
from .models import AuditLog


def _parse_jalali_date(value):
    """تبدیل تاریخ شمسی به میلادی — None اگه نامعتبر بود"""
    if not value:
        return None
    try:
        persian = '۰۱۲۳۴۵۶۷۸۹'
        arabic = '٠١٢٣٤٥٦٧٨٩'
        for i, (p, a) in enumerate(zip(persian, arabic)):
            value = value.replace(p, str(i)).replace(a, str(i))
        value = value.replace('-', '/').replace('.', '/')
        if ' ' in value:
            value = value.split(' ', 1)[0]
        y, m, d = map(int, value.split('/'))
        return jdatetime.date(y, m, d).togregorian()
    except (ValueError, AttributeError):
        return None


def _apply_audit_filters(qs, request):
    """اعمال همه‌ی فیلترها روی queryset (مشترک بین view و export)"""
    action = request.GET.get('action', '').strip()
    if action:
        qs = qs.filter(action=action)

    model_name = request.GET.get('model', '').strip()
    if model_name:
        qs = qs.filter(model_name=model_name)

    user_id = request.GET.get('user', '').strip()
    if user_id:
        qs = qs.filter(user_id=user_id)

    search = request.GET.get('q', '').strip()
    if search:
        qs = qs.filter(
            Q(object_repr__icontains=search) |
            Q(description__icontains=search)
        )

    from_date_str = request.GET.get('from_date', '').strip()
    to_date_str = request.GET.get('to_date', '').strip()
    from_date = _parse_jalali_date(from_date_str)
    to_date = _parse_jalali_date(to_date_str)
    if from_date:
        qs = qs.filter(created_at__date__gte=from_date)
    if to_date:
        qs = qs.filter(created_at__date__lte=to_date)

    return qs


def _get_charts_data(logs_qs):
    """داده‌های نمودارها از queryset فیلترشده"""
    today = timezone.now().date()

    # ─── فعالیت روزانه ۳۰ روز اخیر ───
    thirty_days_ago = today - timedelta(days=29)
    daily_raw = (
        logs_qs.filter(created_at__date__gte=thirty_days_ago)
        .values('created_at__date')
        .annotate(count=Count('id'))
    )
    daily_map = {item['created_at__date']: item['count'] for item in daily_raw}

    daily_labels = []
    daily_data = []
    for i in range(29, -1, -1):
        d = today - timedelta(days=i)
        daily_labels.append(d.strftime('%m/%d'))
        daily_data.append(daily_map.get(d, 0))

    # ─── توزیع اکشن‌ها ───
    action_map = dict(AuditLog.ACTION_CHOICES)
    action_raw = {
        item['action']: item['count']
        for item in logs_qs.values('action').annotate(count=Count('id'))
    }
    action_labels = []
    action_data = []
    for key, label in AuditLog.ACTION_CHOICES:
        if key in action_raw:
            action_labels.append(label)
            action_data.append(action_raw[key])

    # ─── Top 10 کاربران ───
    top_users_raw = (
        logs_qs.exclude(user__isnull=True)
        .values('user__username', 'user__first_name', 'user__last_name')
        .annotate(count=Count('id'))
        .order_by('-count')[:10]
    )
    top_users = []
    for u in top_users_raw:
        name = f"{u['user__first_name'] or ''} {u['user__last_name'] or ''}".strip()
        if not name:
            name = u['user__username']
        top_users.append({
            'name': name,
            'username': u['user__username'],
            'count': u['count'],
        })

    # ─── Top 10 مدل‌ها ───
    top_models = list(
        logs_qs.values('verbose_name')
        .annotate(count=Count('id'))
        .order_by('-count')[:10]
    )

    return {
        'daily_labels': daily_labels,
        'daily_data': daily_data,
        'action_labels': action_labels,
        'action_data': action_data,
        'top_users': top_users,
        'top_models': top_models,
    }


# ═══════════════════════════════════════════════════════════
# Views
# ═══════════════════════════════════════════════════════════

@login_required
@require_permission('audit.view')
def audit_log_list(request):
    """لیست لاگ‌های سیستم با فیلتر + نمودارها"""
    logs = AuditLog.objects.select_related('user', 'content_type').all()
    logs = _apply_audit_filters(logs, request)

    # ─── آمار کلی (unfiltered) ───
    today = jdatetime.date.today().togregorian()
    stats = {
        'total': AuditLog.objects.count(),
        'today': AuditLog.objects.filter(created_at__date=today).count(),
        'created': AuditLog.objects.filter(action='created').count(),
        'updated': AuditLog.objects.filter(action='updated').count(),
        'deleted': AuditLog.objects.filter(action='deleted').count(),
    }

    # ─── چک فیلتر فعال ───
    has_filters = any([
        request.GET.get('action'),
        request.GET.get('model'),
        request.GET.get('user'),
        request.GET.get('q'),
        request.GET.get('from_date'),
        request.GET.get('to_date'),
    ])

    # ─── نمودارها (filtered) ───
    charts = _get_charts_data(logs)

    # ─── dropdowns ───
    models_list = (
        AuditLog.objects.values_list('model_name', 'verbose_name')
        .distinct().order_by('verbose_name')
    )
    users_list = User.objects.filter(
        audit_logs__isnull=False
    ).distinct().order_by('first_name', 'last_name', 'username')

    paginator = Paginator(logs, 50)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    def _json(obj):
        return json.dumps(obj, ensure_ascii=False)

    context = {
        'page_obj': page_obj,
        'stats': stats,
        'action_choices': AuditLog.ACTION_CHOICES,
        'models_list': models_list,
        'users_list': users_list,
        'current_action': request.GET.get('action', '').strip(),
        'current_model': request.GET.get('model', '').strip(),
        'current_user': request.GET.get('user', '').strip(),
        'search': request.GET.get('q', '').strip(),
        'from_date': request.GET.get('from_date', '').strip(),
        'to_date': request.GET.get('to_date', '').strip(),
        'has_filters': has_filters,
        'filtered_count': logs.count() if has_filters else None,

        # Charts
        'top_users': charts['top_users'],
        'top_models': charts['top_models'],
        'daily_labels_json': _json(charts['daily_labels']),
        'daily_data_json': _json(charts['daily_data']),
        'action_labels_json': _json(charts['action_labels']),
        'action_data_json': _json(charts['action_data']),
    }
    return render(request, 'audit/log_list.html', context)


@login_required
@require_permission('audit.view')
def audit_export_excel(request):
    """خروجی Excel از لاگ‌های فیلترشده"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        return HttpResponse('openpyxl نصب نیست.', status=500)

    logs = AuditLog.objects.select_related('user').all()
    logs = _apply_audit_filters(logs, request)

    wb = Workbook()
    ws = wb.active
    ws.title = 'لاگ سیستم'
    ws.sheet_view.rightToLeft = True

    header_font = Font(bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill('solid', fgColor='2E86AB')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center')
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    # ─── عنوان ───
    ws.append(['گزارش لاگ سیستم — CRM دیجیران'])
    ws.merge_cells('A1:G1')
    ws['A1'].font = Font(bold=True, size=14, color='2E86AB')
    ws['A1'].alignment = center
    ws.row_dimensions[1].height = 28

    ws.append([])
    ws.append([f'تعداد رکورد: {logs.count()}'])
    ws.append([])

    # ─── هدر جدول ───
    headers = ['تاریخ و ساعت', 'کاربر', 'مدل', 'شیء', 'رویداد', 'توضیحات', 'IP']
    ws.append(headers)
    header_row = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=header_row, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    # ─── داده‌ها (سقف ۵۰۰۰ رکورد) ───
    for log in logs[:5000]:
        ws.append([
            log.created_at.strftime('%Y/%m/%d %H:%M'),
            log.user_display,
            log.verbose_name,
            log.object_repr,
            log.get_action_display(),
            (log.description or '')[:200],
            log.ip_address or '—',
        ])
        for c in ws[ws.max_row]:
            c.alignment = right
            c.border = border

    # ─── عرض ستون‌ها ───
    widths = [18, 18, 20, 30, 12, 50, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = w

    # ─── پاسخ ───
    import io
    from datetime import datetime as _dt
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    filename = f'audit-log-{_dt.now().strftime("%Y%m%d-%H%M")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response